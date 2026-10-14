#!/usr/bin/env python3
"""Independent six-field atlas validation; write aggregate metadata only.

This command has no patient export or approval mode. It never changes the
workbook, original adjudication records, crosswalk, or existing export policy.
Exit 0 means every scoped prerequisite passed, not permission to preview or
publish. Exit 1 means the core validation remains blocked.
"""

import argparse
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import stat
import subprocess
from typing import Any
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

import openpyxl

from normalization import REVIEW_RECORD_KEYS, SourceError, normalize_core_row


REPOSITORY = Path(__file__).resolve().parents[2]
SPEC_VERSION = "1.0.0"
SCHEMA_VERSION = 1
MASTER_SHEET = "Master_Patient_Dataset_N38"
DEFAULT_WORKBOOK = REPOSITORY / "esearch_sources/Triple_M_Master_Patient_Dataset_Expanded_N38.xlsx"
FROZEN_SOURCE_SHA256 = "d6ced420d83eb046559ea445422ea7a838990989ab535a2f4530ef40db674fcb"
PRIVATE_DIRECTORY = REPOSITORY / "esearch_sources/_processed"
DEFAULT_CROSSWALK = PRIVATE_DIRECTORY / "atlas.provenance.private.json"
DEFAULT_SIU_REVIEW = PRIVATE_DIRECTORY / "core_atlas/siu_provenance_review.private.json"
DEFAULT_FULL_AUDIT = PRIVATE_DIRECTORY / "validation.audit.private.json"
FULL_AUDIT_CHECK_COUNT = 48  # Existing full-workflow contract; not a core prerequisite.
FULL_AUDIT_REQUIRED_NAMES = frozenset({
    "Supported quantitative peak troponin n", "Troponin median (ng/L)",
    "Troponin IQR (ng/L)", "Troponin range (ng/L, audit)",
    "Original workbook unchanged during ingestion",
})
CORE_HEADERS = (
    "Patient_ID", "DOI", "RNS_Performed", "RNS_Result", "SFEMG_Performed",
    "SFEMG_Result", "EMG_Conventional_Performed", "EMG_Interpretation",
    "Objective_NMJ_Category", "AChR_Performed", "AChR_Positive", "Outcome_Survival",
)
AVAILABILITY_VALUES = frozenset({"performed", "not_performed", "not_reported", "not_documented", "unknown"})
FORMAL_RESULT_VALUES = frozenset({"negative", "abnormal", "not_reported", "not_documented"})
BROWSER_VALUES = {
    "achr": frozenset({"positive", "negative", "not_documented"}),
    "formalNmj": frozenset({"documented", "not_documented"}),
    "nmjEvidence": frozenset({"formal_negative", "indeterminate_qualitative", "not_documented"}),
    "needleEmg": frozenset({"myopathic", "neuropathic", "normal", "indeterminate", "not_documented"}),
    "vitalStatus": frozenset({"death", "survival", "unknown"}),
}


@dataclass(frozen=True)
class CoreCheck:
    name: str
    expected: Any
    observed: Any
    passed: bool
    domain: str
    note: str = ""


def comparison(name: str, expected: Any, observed: Any, domain: str, note: str = "") -> CoreCheck:
    return CoreCheck(name, expected, observed, expected == observed, domain, note)


def sha256(path: Path) -> str | None:
    try:
        return hashlib.sha256(path.read_bytes()).hexdigest()
    except OSError:
        return None


def read_json(path: Path) -> dict | None:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
        return value if isinstance(value, dict) else None
    except (OSError, ValueError):
        return None


def load_core_source(workbook: Path) -> tuple[list[dict], list[CoreCheck]]:
    """Read only the required field columns; reject errors/formulas before mapping.

    All physical data rows are processed, including incomplete/empty rows, so
    normalization cannot silently drop them. Extra source columns neither
    influence category parsing nor become downgraded full-study failures.
    """
    checks: list[CoreCheck] = []
    values = formulas = None
    try:
        values = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
        formulas = openpyxl.load_workbook(workbook, read_only=True, data_only=False)
        has_sheet = MASTER_SHEET in values.sheetnames and MASTER_SHEET in formulas.sheetnames
        checks.append(comparison("Authoritative master sheet available", True, has_sheet, "integrity"))
        if not has_sheet:
            return [], checks
        sheet, formula_sheet = values[MASTER_SHEET], formulas[MASTER_SHEET]
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        invalid_count = sum(header.count(name) != 1 for name in CORE_HEADERS)
        checks.append(comparison("Every scoped header maps exactly once", 0, invalid_count, "integrity"))
        if invalid_count:
            return [], checks
        columns = {name: header.index(name) + 1 for name in CORE_HEADERS}
        row_count = max(sheet.max_row - 1, 0)
        checks.append(comparison("All physical master data rows", 38, row_count, "integrity"))
        source_rows = {number: {} for number in range(2, sheet.max_row + 1)}
        formula_count = error_count = 0
        for name, column in columns.items():
            for number, cells in enumerate(sheet.iter_rows(min_row=2, min_col=column, max_col=column, values_only=True), 2):
                source_rows[number][name] = cells[0]
            for cells in formula_sheet.iter_rows(min_row=2, min_col=column, max_col=column):
                formula_count += cells[0].data_type == "f"
                error_count += cells[0].data_type == "e"
        checks.append(comparison("Scoped input formula cells", 0, formula_count, "integrity", "Cached formula values are not accepted"))
        checks.append(comparison("Scoped input spreadsheet error cells", 0, error_count, "integrity"))
        if formula_count or error_count:
            return [], checks
        records, error_count = [], 0
        for number, row in source_rows.items():
            try:
                records.append(normalize_core_row(row, number))
            except (SourceError, KeyError, TypeError, ValueError):
                error_count += 1
        checks.append(comparison("Scoped normalization errors without inference", 0, error_count, "integrity"))
        checks.append(comparison("No source rows dropped during scoped mapping", row_count, len(records), "integrity"))
        return records, checks
    except (OSError, ValueError, KeyError, BadZipFile, ParseError, openpyxl.utils.exceptions.InvalidFileException):
        checks.append(comparison("Authoritative workbook readable", True, False, "integrity"))
        return [], checks
    finally:
        if values is not None:
            values.close()
        if formulas is not None:
            formulas.close()


def validate_core_records(records: list[dict]) -> list[CoreCheck]:
    """Own versioned core rules, independent of full-study check selection."""
    checks: list[CoreCheck] = []

    def check(name: str, expected: Any, observed: Any, domain: str = "numerical", note: str = "") -> None:
        checks.append(comparison(name, expected, observed, domain, note))

    def counts(field: str, categories: tuple[str, ...]) -> dict:
        observed = Counter(row[field] for row in records)
        return {category: observed[category] for category in categories}

    check("Actual normalized patient records", 38, len(records))
    check("Unique private patient identifiers", 38, len({r["source_id"] for r in records}), "integrity")
    check("Unique canonical source publications", 27, len({r["source_doi"] for r in records}))
    check("Private source rows map one-to-one", True, len({r["source_row"] for r in records}) == len(records), "integrity")
    valid_source = all(isinstance(r["source_id"], str) and bool(r["source_id"]) and
                       isinstance(r["source_doi"], str) and bool(r["source_doi"]) and
                       isinstance(r["source_row"], int) and not isinstance(r["source_row"], bool) and r["source_row"] >= 2 for r in records)
    check("Private source identity fields present and typed", True, valid_source, "integrity")
    valid_categories = all(all(r[f"{test}_status"] in AVAILABILITY_VALUES for test in ("rns", "sfemg", "achr", "needle")) and
                           all(r[f"{test}_result"] in FORMAL_RESULT_VALUES for test in ("rns", "sfemg")) and
                           r["achr_result"] in {"positive", "negative", "not_reported", "not_documented", "unknown"} and
                           r["nmj_evidence"] in BROWSER_VALUES["nmjEvidence"] and
                           r["needle_emg"] in BROWSER_VALUES["needleEmg"] | {"not_reported", "not_performed"} and
                           r["vital_status"] in BROWSER_VALUES["vitalStatus"] and isinstance(r["formal_nmj"], bool) for r in records)
    check("Every internal core category has an allowed value", True, valid_categories, "integrity")
    tested = [r for r in records if r["achr_status"] == "performed"]
    positive = [r for r in records if r["achr_result"] == "positive"]
    negative = [r for r in records if r["achr_result"] == "negative"]
    check("AChR tested / positive / negative / not documented", [30, 11, 19, 8], [len(tested), len(positive), len(negative), len(records) - len(tested)])
    check("AChR availability agrees with documented result", 0, sum((r["achr_status"] == "performed") != (r["achr_result"] in {"positive", "negative"}) for r in records))
    rns = {r["source_id"] for r in records if r["rns_status"] == "performed"}
    sfemg = {r["source_id"] for r in records if r["sfemg_status"] == "performed"}
    formal = {r["source_id"] for r in records if r["formal_nmj"]}
    check("Unique RNS and/or SFEMG patients", 4, len(rns | sfemg))
    check("RNS evaluated patients", 4, len(rns))
    check("SFEMG evaluated patients", 1, len(sfemg))
    check("SFEMG is nested in the RNS cohort", True, sfemg <= rns)
    check("Formal availability matches each row's RNS/SFEMG union", True, all(r["formal_nmj"] == (r["rns_status"] == "performed" or r["sfemg_status"] == "performed") for r in records))
    check("Formal patient identity set equals testing union", True, formal == rns | sfemg)
    performed = [(r, test) for r in records for test in ("rns", "sfemg") if r[f"{test}_status"] == "performed"]
    check("Performed formal studies have explicit negative results", True, all(r[f"{test}_result"] == "negative" for r, test in performed))
    abnormal = {r["source_id"] for r, test in performed if r[f"{test}_result"] == "abnormal"}
    check("Abnormal formal NMJ patients / evaluated patients", [0, 4], [len(abnormal), len(rns | sfemg)])
    check("Unperformed formal studies never gain a result", 0, sum(r[f"{test}_status"] != "performed" and r[f"{test}_result"] in {"negative", "abnormal"} for r in records for test in ("rns", "sfemg")))
    check("Objective NMJ evidence categories", {"formal_negative": 4, "indeterminate_qualitative": 2, "not_documented": 32}, counts("nmj_evidence", ("formal_negative", "indeterminate_qualitative", "not_documented")))
    check("Formal-negative evidence requires explicit formal negativity", True, all((r["nmj_evidence"] == "formal_negative") == r["formal_nmj"] for r in records))
    check("Qualitative indeterminate evidence has no formal confirmation", 0, sum(r["nmj_evidence"] == "indeterminate_qualitative" and r["formal_nmj"] for r in records))
    check("AChR-positive and formally evaluated overlap", 0, len({r["source_id"] for r in positive} & (rns | sfemg)), note="Unpaired evidence is not negative physiological evidence")
    needle = [r for r in records if r["needle_status"] == "performed"]
    needle_values = ("myopathic", "neuropathic", "normal", "indeterminate")
    check("Needle EMG evaluated / not documented", [11, 27], [len(needle), len(records) - len(needle)])
    check("Needle EMG categories among evaluated", {"myopathic": 6, "neuropathic": 1, "normal": 1, "indeterminate": 3}, {category: sum(r["needle_emg"] == category for r in needle) for category in needle_values})
    check("Needle availability agrees with an interpretation", 0, sum((r["needle_status"] == "performed") != (r["needle_emg"] in needle_values) for r in records))
    check("Documented vital status categories", {"death": 16, "survival": 21, "unknown": 1}, counts("vital_status", ("death", "survival", "unknown")))
    check("Unknown vital status remains outside known denominator", 37, sum(r["vital_status"] != "unknown" for r in records))
    return checks


def private_artifact_safe(path: Path) -> bool:
    """Read-only check: private artifacts must be ignored, untracked and 0600."""
    try:
        if path.is_symlink() or not path.is_file() or not path.resolve().is_relative_to(PRIVATE_DIRECTORY.resolve()):
            return False
        mode = stat.S_IMODE(path.stat().st_mode)
        ignored = subprocess.run(["git", "check-ignore", "-q", str(path)], cwd=REPOSITORY, check=False, capture_output=True)
        tracked = subprocess.run(["git", "ls-files", "--", str(path)], cwd=REPOSITORY, text=True, check=True, capture_output=True)
        return mode == 0o600 and ignored.returncode == 0 and not tracked.stdout.strip()
    except (OSError, ValueError, subprocess.SubprocessError):
        return False


def validate_crosswalk(records: list[dict], payload: dict | None, source_hash: str | None) -> list[CoreCheck]:
    valid_hash = bool(payload and source_hash and payload.get("sourceSha256") == source_hash and payload.get("schemaVersion") == 1)
    entries = payload.get("records", []) if payload else []
    expected = [{"label": f"Case {index:02d}", "sourceId": r["source_id"], "sourceRow": r["source_row"], "sourceDoi": r["source_doi"]}
                for index, r in enumerate(sorted(records, key=lambda r: r["source_id"]), 1)]
    mapped = bool(records) and len({r["source_id"] for r in records}) == len(records) and isinstance(entries, list) and entries == expected
    return [comparison("Existing private crosswalk is versioned and source-bound", True, valid_hash, "provenance"),
            comparison("Existing labels map actual source identity and rows exactly", True, mapped, "provenance", "Original crosswalk is read, never rewritten; no patient records are serialized")]


def validate_siu_provenance(records: list[dict], payload: dict | None, source_hash: str | None) -> list[CoreCheck]:
    target = next((r for r in records if r["source_row"] == 39), None)
    bound = bool(payload and target and source_hash and payload.get("schemaVersion") == 1 and
                 payload.get("sourceSha256") == source_hash and payload.get("sourceId") == target["source_id"] and
                 payload.get("sourceRow") == target["source_row"] and isinstance(payload.get("sourceDoi"), str) and
                 payload["sourceDoi"].strip().lower() == target["source_doi"])
    linked_target = next((r for r in records if r["source_row"] == 22), None)
    linked_expected = [{"sourceRow": linked_target["source_row"], "sourceId": linked_target["source_id"], "sourceDoi": linked_target["source_doi"]}] if linked_target else []
    linked_bound = bool(payload and linked_target and payload.get("linkedSourceRows") == [22] and
                        payload.get("linkedSourceMappings") == linked_expected)
    evidence = payload.get("evidence", []) if payload else []
    structured = isinstance(evidence, list) and bool(evidence) and all(isinstance(item, dict) and
                 all(item.get(key) for key in ("url", "location", "facts")) for item in evidence)
    linked_resolved = bool(payload and linked_bound and payload.get("linkedSourceMappingsVerified") is True and
                           payload.get("potentialDuplicateResolved") is True and payload.get("categoricalDisagreementsResolved") is True)
    resolved = bool(payload and bound and structured and linked_resolved and payload.get("status") == "RESOLVED" and
                    payload.get("mappingVerified") is True and payload.get("membershipChangeProposed") is False and
                    payload.get("authorDecisionRequired") is False and payload.get("questions") == [])
    return [comparison("Targeted Siu provenance review bound to actual frozen row", True, bound, "provenance"),
            comparison("Targeted review has evidence locations and facts", True, structured, "provenance"),
            comparison("Linked source review bound to both actual frozen mappings", True, linked_bound, "provenance"),
            comparison("Linked source identity, possible duplication and category disagreements resolved", True, linked_resolved, "provenance", "Frozen aggregate agreement cannot resolve clinical case identity or contradictory source categories"),
            comparison("Targeted Siu source mapping resolved without unauthorized membership change", True, resolved, "provenance", "An unresolved source mapping prevents CORE_ATLAS_VALIDATION PASS even when numerical checks pass")]


def full_study_summary(audit: dict | None, source_hash: str | None) -> dict:
    """Separate full-study status from its own executed audit, never a core rule."""
    checks = audit.get("checks", []) if audit else []
    valid = bool(audit and source_hash and audit.get("sourceSha256") == source_hash and
                 isinstance(audit.get("validationPassed"), bool) and isinstance(checks, list) and
                 len(checks) == FULL_AUDIT_CHECK_COUNT and all(isinstance(check, dict) and
                 isinstance(check.get("name"), str) and isinstance(check.get("passed"), bool) for check in checks))
    if valid:
        names = {check["name"] for check in checks}
        valid = len(names) == len(checks) and FULL_AUDIT_REQUIRED_NAMES <= names
    if not valid:
        return {"status": "FAIL", "auditVerified": False, "checksPassed": None, "checksTotal": None, "failedChecks": []}
    passed = all(check["passed"] for check in checks) and audit.get("validationPassed") is True
    return {"status": "PASS" if passed else "FAIL", "auditVerified": True,
            "checksPassed": sum(check["passed"] for check in checks), "checksTotal": len(checks),
            "failedChecks": [check["name"] for check in checks if not check["passed"]]}


def status_payload(checks: list[CoreCheck], source_hash: str | None, timestamp: str, full_summary: dict) -> dict:
    numerical = [check for check in checks if check.domain == "numerical"]
    core_passed = bool(numerical) and bool(checks) and all(check.passed for check in checks)
    return {
        "schemaVersion": SCHEMA_VERSION, "specVersion": SPEC_VERSION, "validatedAt": timestamp,
        "sourceSha256": source_hash, "dataScope": "six-field-neuromuscular-atlas",
        "statuses": {
            "FULL_STUDY_VALIDATION": full_summary["status"],
            "CORE_ATLAS_NUMERICAL_VALIDATION": "PASS" if numerical and all(check.passed for check in numerical) else "FAIL",
            "CORE_ATLAS_VALIDATION": "PASS" if core_passed else "FAIL",
            "CORE_ATLAS_PUBLICATION_APPROVAL": "PENDING",
            "CORE_ATLAS_PRIVATE_PREVIEW_APPROVAL": "PENDING",
            "CORE_ATLAS_EXPORT_POLICY": "NOT_ACTIVATED",
        },
        "fullStudyAudit": full_summary,
        "checksPassed": sum(check.passed for check in checks), "checksTotal": len(checks),
        "checks": [asdict(check) for check in checks],
        "patientExportGenerated": False, "privatePreviewEnabled": False, "publicReleaseEnabled": False,
    }


def render_core_report(payload: dict) -> str:
    def cell(value: Any) -> str:
        return json.dumps(value, ensure_ascii=False).replace("|", "\\|")

    lines = ["# Independent core atlas validation report", "", f"Executed at: {payload['validatedAt']}.", "",
             f"Specification: `CORE_NEUROMUSCULAR_ATLAS_VALIDATION.md`, version {payload['specVersion']}.", "",
             "This is an independent scoped workflow. Numerical agreement does not resolve source provenance, validate the full manuscript, or authorize disclosure.", "",
             "| Status | Result |", "| --- | --- |"]
    lines += [f"| {name} | {value} |" for name, value in payload["statuses"].items()]
    full = payload["fullStudyAudit"]
    lines += ["", f"Scoped prerequisite checks: {payload['checksPassed']}/{payload['checksTotal']} passed.", "",
              f"Full-study audit verified against this source: {cell(full['auditVerified'])}; required checks passed/total: {cell([full['checksPassed'], full['checksTotal']])}.", "",
              "| Scoped check | Domain | Expected | Observed | Result |", "| --- | --- | --- | --- | --- |"]
    lines += [f"| {check['name']} | {check['domain']} | {cell(check['expected'])} | {cell(check['observed'])} | {'PASS' if check['passed'] else 'FAIL'} |" for check in payload["checks"]]
    lines += ["", "The required 38 records, 27 publication keys, categorical denominators and testing unions are evaluated from the actual frozen rows. No rows are reconstructed from aggregate targets. Missing or unperformed testing is never a negative result.", "",
              "The Siu review remains a mandatory source-integrity prerequisite. The private review must establish both linked case mappings, resolve the possible source misattribution or duplicate and reconcile the category disagreements without an unauthorized membership change. Unresolved questions remain blocking failures. Numerical reproduction of source rows is not proof of 38 clinically distinct patients or 27 correctly attributed publications.", "",
              "The full-study workflow retains all frozen targets and its failed troponin checks. Its separate audit is summarized for status reporting; no full-study failure is filtered, converted to a warning, or used to define this specification.", "",
              "No patient-level export was generated, loaded, approved, committed or published by this workflow. It does not write or replace any original private crosswalk or adjudication record. The existing export policy remains active; a core export policy has not been activated. Private preview and publication require separate explicit author authorization of concrete, exact records under the documented privacy review.", "",
              "The JSON companion contains validation metadata and aggregate checks only. It contains no browser records, raw source identifiers, DOI crosswalk, source row lists, clinical narratives, or review evidence. Patient UI interactions cannot be claimed as tested by this validation."]
    return "\n".join(lines) + "\n"


def run_core_validation(workbook: Path, report: Path, status: Path,
                        crosswalk: Path = DEFAULT_CROSSWALK, siu_review: Path = DEFAULT_SIU_REVIEW,
                        full_audit: Path = DEFAULT_FULL_AUDIT) -> dict:
    # Aggregate outputs cannot overwrite any immutable input or research artifact.
    protected = {path.resolve() for path in (workbook, crosswalk, siu_review, full_audit)}
    research_roots = (REPOSITORY / "esearch_sources", REPOSITORY / "research_sources")
    for destination in (report, status):
        if destination.resolve() in protected or any(destination.resolve().is_relative_to(root.resolve()) for root in research_roots):
            raise ValueError("Aggregate reports must not overwrite research sources or private provenance")
    if report.resolve() == status.resolve():
        raise ValueError("Report and status require distinct output files")
    timestamp = datetime.now(timezone.utc).isoformat(timespec="seconds")
    source_hash = sha256(workbook)
    records, checks = load_core_source(workbook)
    checks += validate_core_records(records)
    checks += [comparison("Frozen authoritative source fingerprint", FROZEN_SOURCE_SHA256, source_hash, "integrity"),
               comparison("Original workbook unchanged during scoped ingestion", True, bool(source_hash and sha256(workbook) == source_hash), "integrity")]
    checks += validate_crosswalk(records, read_json(crosswalk), source_hash)
    checks += validate_siu_provenance(records, read_json(siu_review), source_hash)
    checks += [comparison("Private crosswalk remains ignored, untracked and owner-only", True, private_artifact_safe(crosswalk), "privacy"),
               comparison("Private targeted provenance remains ignored, untracked and owner-only", True, private_artifact_safe(siu_review), "privacy"),
               comparison("Existing patient review export remains absent", False, (PRIVATE_DIRECTORY / "atlas.review.json").exists(), "privacy"),
               comparison("Browser field whitelist remains exactly six scoped fields", {"label", *BROWSER_VALUES}, set(REVIEW_RECORD_KEYS), "privacy")]
    # Replace sets with stable arrays before any JSON serialization.
    checks = [CoreCheck(check.name, sorted(check.expected) if isinstance(check.expected, set) else check.expected,
                        sorted(check.observed) if isinstance(check.observed, set) else check.observed,
                        check.passed, check.domain, check.note) for check in checks]
    payload = status_payload(checks, source_hash, timestamp, full_study_summary(read_json(full_audit), source_hash))
    report.parent.mkdir(parents=True, exist_ok=True)
    status.parent.mkdir(parents=True, exist_ok=True)
    report.write_text(render_core_report(payload), encoding="utf-8")
    status.write_text(json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    return payload


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--report", type=Path, default=REPOSITORY / "docs/CORE_ATLAS_VALIDATION_REPORT.md")
    parser.add_argument("--status", type=Path, default=REPOSITORY / "docs/CORE_ATLAS_VALIDATION_STATUS.json")
    parser.add_argument("--full-audit", type=Path, default=DEFAULT_FULL_AUDIT, help="Source-bound full-study aggregate audit; does not define core rules")
    arguments = parser.parse_args()
    try:
        payload = run_core_validation(arguments.workbook.resolve(), arguments.report.resolve(), arguments.status.resolve(), full_audit=arguments.full_audit.resolve())
    except (OSError, ValueError, subprocess.SubprocessError) as exc:
        parser.exit(1, f"Core validation report blocked: {type(exc).__name__}\n")
    print(json.dumps(payload["statuses"], sort_keys=True))
    print(f"{payload['checksPassed']}/{payload['checksTotal']} scoped checks passed. Aggregate metadata only; no patient export or release authorization.")
    return 0 if payload["statuses"]["CORE_ATLAS_VALIDATION"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
