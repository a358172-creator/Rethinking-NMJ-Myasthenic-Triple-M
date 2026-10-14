"""Independent core regressions; temporary fixtures never enter a browser."""

from copy import deepcopy
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from core_validation import (  # noqa: E402
    CORE_HEADERS, DEFAULT_CROSSWALK, DEFAULT_FULL_AUDIT, DEFAULT_SIU_REVIEW,
    DEFAULT_WORKBOOK, FROZEN_SOURCE_SHA256, FULL_AUDIT_CHECK_COUNT,
    FULL_AUDIT_REQUIRED_NAMES, MASTER_SHEET, comparison,
    full_study_summary, load_core_source, private_artifact_safe,
    run_core_validation, sha256, status_payload, validate_core_records,
    validate_crosswalk, validate_siu_provenance,
)
from normalization import (  # noqa: E402
    SourceError, achr_result, availability, negative_result, normalize_core_row,
    normalize_row, review_records,
)
from process_workbook import load_source, validate_records  # noqa: E402


def raw_core_fixture() -> dict:
    return {
        "Patient_ID": "unit-test-only", "DOI": "unit-test-only",
        "RNS_Performed": "Not reported", "RNS_Result": "Not reported",
        "SFEMG_Performed": "Not reported", "SFEMG_Result": "Not reported",
        "EMG_Conventional_Performed": "Not reported", "EMG_Interpretation": "Not reported",
        "Objective_NMJ_Category": "D. NOT TESTED / NOT REPORTED",
        "AChR_Performed": "Not reported", "AChR_Positive": "Not reported",
        "Outcome_Survival": "Unknown",
    }


class ScopedMappingTests(unittest.TestCase):
    def test_missingness_remains_distinct_and_never_negative(self):
        for source, expected in (("No", "not_performed"), ("Not reported", "not_reported"), (None, "not_documented"), ("Unknown", "unknown")):
            self.assertEqual(availability(source), expected)
            if source != "No":
                self.assertNotEqual(achr_result(source), "negative")
        self.assertEqual(negative_result(None), "not_documented")
        self.assertEqual(negative_result("Not reported"), "not_reported")
        normalized = normalize_core_row(raw_core_fixture(), 2)
        browser, _ = review_records([normalized])
        self.assertEqual(browser[0]["achr"], "not_documented")
        self.assertEqual(browser[0]["formalNmj"], "not_documented")
        self.assertEqual(browser[0]["nmjEvidence"], "not_documented")
        self.assertEqual(browser[0]["needleEmg"], "not_documented")
        self.assertEqual(browser[0]["vitalStatus"], "unknown")

    def test_core_never_parses_excluded_biochemical_values(self):
        row = {**raw_core_fixture(), "Troponin_Max": {"invalid": "full-only"},
               "Age": "Late 70s", "LVEF_Reduced": "not a mapped category", "MG_Phenotype": "Yes"}
        core = normalize_core_row(row, 2)
        self.assertNotIn("troponin", core)
        self.assertNotIn("age", core)
        self.assertNotIn("lvef", core)
        with self.assertRaises(SourceError):
            normalize_row(row, 2)

    def test_loader_accepts_only_core_headers_and_ignores_unrelated_formula(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "scoped-test.xlsx"
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = MASTER_SHEET
            sheet.append([*CORE_HEADERS, "Troponin_Max"])
            row = raw_core_fixture()
            sheet.append([row[name] for name in CORE_HEADERS] + ["=1/0"])
            book.save(path)
            book.close()
            records, checks = load_core_source(path)
            self.assertEqual(len(records), 1)
            self.assertTrue(next(c.passed for c in checks if c.name == "Scoped input formula cells"))
            self.assertTrue(next(c.passed for c in checks if c.name == "Scoped normalization errors without inference"))
            self.assertFalse(any("troponin" in c.name.lower() or "warning" in c.domain.lower() for c in checks))
            _, full_checks, _, _ = load_source(path)
            self.assertFalse(all(c.passed for c in full_checks))

    def test_scoped_formula_and_spreadsheet_error_block_ingestion(self):
        for value in ("=1+1", "#DIV/0!"):
            with self.subTest(value=value), tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "invalid-test.xlsx"
                book = openpyxl.Workbook()
                sheet = book.active
                sheet.title = MASTER_SHEET
                sheet.append(CORE_HEADERS)
                row = raw_core_fixture()
                row["AChR_Positive"] = value
                sheet.append([row[name] for name in CORE_HEADERS])
                book.save(path)
                book.close()
                records, checks = load_core_source(path)
                self.assertEqual(records, [])
                selected = [c for c in checks if c.name in {"Scoped input formula cells", "Scoped input spreadsheet error cells"}]
                self.assertEqual(sum(not c.passed for c in selected), 1)

    def test_required_header_duplication_blocks_core_mapping(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "duplicate-test.xlsx"
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = MASTER_SHEET
            sheet.append([*CORE_HEADERS, "AChR_Positive"])
            book.save(path)
            book.close()
            records, checks = load_core_source(path)
            self.assertEqual(records, [])
            self.assertFalse(next(c.passed for c in checks if c.name == "Every scoped header maps exactly once"))

    def test_missing_workbook_fails_closed(self):
        records, checks = load_core_source(Path("/nonexistent/core-workbook.xlsx"))
        self.assertEqual(records, [])
        self.assertFalse(all(c.passed for c in checks))

    def test_private_artifact_missing_or_symlink_fails(self):
        self.assertFalse(private_artifact_safe(Path("/nonexistent/provenance.private.json")))
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "linked.private.json"
            target = Path(directory) / "target.json"
            target.write_text("{}", encoding="utf-8")
            path.symlink_to(target)
            self.assertFalse(private_artifact_safe(path))


class SeparateAuthorizationTests(unittest.TestCase):
    def test_numerical_and_core_pass_never_authorize_preview_or_publication(self):
        payload = status_payload([comparison("scoped", True, True, "numerical")], FROZEN_SOURCE_SHA256,
                                 "test-only", {"status": "FAIL"})
        self.assertEqual(payload["statuses"]["CORE_ATLAS_VALIDATION"], "PASS")
        self.assertEqual(payload["statuses"]["FULL_STUDY_VALIDATION"], "FAIL")
        self.assertEqual(payload["statuses"]["CORE_ATLAS_PUBLICATION_APPROVAL"], "PENDING")
        self.assertEqual(payload["statuses"]["CORE_ATLAS_PRIVATE_PREVIEW_APPROVAL"], "PENDING")
        self.assertEqual(payload["statuses"]["CORE_ATLAS_EXPORT_POLICY"], "NOT_ACTIVATED")
        self.assertFalse(payload["patientExportGenerated"])
        self.assertFalse(payload["privatePreviewEnabled"])
        self.assertFalse(payload["publicReleaseEnabled"])

    def test_provenance_failure_blocks_core_while_numerical_passes(self):
        checks = [comparison("numerical", True, True, "numerical"), comparison("source provenance", True, False, "provenance")]
        payload = status_payload(checks, FROZEN_SOURCE_SHA256, "test-only", {"status": "FAIL"})
        self.assertEqual(payload["statuses"]["CORE_ATLAS_NUMERICAL_VALIDATION"], "PASS")
        self.assertEqual(payload["statuses"]["CORE_ATLAS_VALIDATION"], "FAIL")

    def test_full_status_requires_current_hash_bound_audit(self):
        names = list(FULL_AUDIT_REQUIRED_NAMES) + [f"test-only check {index}" for index in range(FULL_AUDIT_CHECK_COUNT - len(FULL_AUDIT_REQUIRED_NAMES))]
        failed_name = "Troponin median (ng/L)"
        audit = {"sourceSha256": FROZEN_SOURCE_SHA256, "validationPassed": True,
                 "checks": [{"name": name, "passed": name != failed_name} for name in names]}
        summary = full_study_summary(audit, FROZEN_SOURCE_SHA256)
        self.assertEqual(summary["status"], "FAIL")
        self.assertEqual(summary["failedChecks"], [failed_name])
        self.assertFalse(full_study_summary(audit, "different source")["auditVerified"])
        self.assertEqual(full_study_summary(None, FROZEN_SOURCE_SHA256)["status"], "FAIL")
        audit["checks"] = audit["checks"][:1]
        self.assertFalse(full_study_summary(audit, FROZEN_SOURCE_SHA256)["auditVerified"])

    def test_report_destination_cannot_overwrite_immutable_inputs(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            source.write_bytes(b"protected fixture")
            original = source.read_bytes()
            with self.assertRaises(ValueError):
                run_core_validation(source, source, Path(directory) / "status.json")
            self.assertEqual(source.read_bytes(), original)


@unittest.skipUnless(DEFAULT_WORKBOOK.exists(), "Authoritative private workbook is not supplied")
class AuthoritativeCoreTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.records, cls.ingestion_checks = load_core_source(DEFAULT_WORKBOOK)
        cls.checks = validate_core_records(cls.records)

    def check_named(self, checks, name):
        return next(c for c in checks if c.name == name)

    def resolved_review(self):
        target = next(r for r in self.records if r["source_row"] == 39)
        linked = next(r for r in self.records if r["source_row"] == 22)
        return {
            "schemaVersion": 1, "sourceSha256": FROZEN_SOURCE_SHA256,
            "sourceRow": target["source_row"], "sourceId": target["source_id"], "sourceDoi": target["source_doi"],
            "status": "RESOLVED", "mappingVerified": True, "membershipChangeProposed": False,
            "authorDecisionRequired": False, "questions": [],
            "evidence": [{"url": "test-only", "location": "test-only", "facts": ["test-only"]}],
            "linkedSourceRows": [22], "linkedSourceMappings": [{"sourceRow": linked["source_row"], "sourceId": linked["source_id"], "sourceDoi": linked["source_doi"]}],
            "linkedSourceMappingsVerified": True, "potentialDuplicateResolved": True,
            "categoricalDisagreementsResolved": True,
        }

    def test_all_independently_authored_core_numerical_invariants(self):
        self.assertEqual(len(self.records), 38)
        self.assertTrue(all(c.passed for c in self.ingestion_checks + self.checks))
        for name, expected in (
            ("Unique canonical source publications", 27),
            ("AChR tested / positive / negative / not documented", [30, 11, 19, 8]),
            ("Unique RNS and/or SFEMG patients", 4),
            ("SFEMG evaluated patients", 1),
            ("Abnormal formal NMJ patients / evaluated patients", [0, 4]),
            ("Needle EMG evaluated / not documented", [11, 27]),
        ):
            self.assertEqual(self.check_named(self.checks, name).observed, expected)

    def test_formal_union_uses_actual_patient_identity(self):
        changed = deepcopy(self.records)
        original = next(r for r in changed if r["sfemg_status"] == "performed")
        other = next(r for r in changed if not r["formal_nmj"])
        original["sfemg_status"], original["sfemg_result"] = "not_reported", "not_reported"
        other["sfemg_status"], other["sfemg_result"], other["formal_nmj"] = "performed", "negative", True
        checks = validate_core_records(changed)
        self.assertTrue(self.check_named(checks, "SFEMG evaluated patients").passed)
        self.assertFalse(self.check_named(checks, "SFEMG is nested in the RNS cohort").passed)
        self.assertEqual(self.check_named(checks, "Unique RNS and/or SFEMG patients").observed, 5)

    def test_forced_negative_missing_result_fails(self):
        changed = deepcopy(self.records)
        missing = next(r for r in changed if r["achr_status"] != "performed")
        missing["achr_result"] = "negative"
        checks = validate_core_records(changed)
        self.assertFalse(self.check_named(checks, "AChR availability agrees with documented result").passed)
        absent_rns = next(r for r in changed if r["rns_status"] != "performed")
        absent_rns["rns_result"] = "negative"
        self.assertFalse(self.check_named(validate_core_records(changed), "Unperformed formal studies never gain a result").passed)

    def test_achr_positive_formal_overlap_fails(self):
        changed = deepcopy(self.records)
        positive = next(r for r in changed if r["achr_result"] == "positive")
        positive["rns_status"], positive["rns_result"], positive["formal_nmj"] = "performed", "negative", True
        self.assertEqual(self.check_named(validate_core_records(changed), "AChR-positive and formally evaluated overlap").observed, 1)

    def test_duplicate_identifier_and_source_row_fail(self):
        changed = deepcopy(self.records)
        changed[1]["source_id"], changed[1]["source_row"] = changed[0]["source_id"], changed[0]["source_row"]
        checks = validate_core_records(changed)
        self.assertFalse(self.check_named(checks, "Unique private patient identifiers").passed)
        self.assertFalse(self.check_named(checks, "Private source rows map one-to-one").passed)

    def test_crosswalk_mappings_and_fingerprint_must_match_exactly(self):
        _, crosswalk = review_records(self.records)
        payload = {"schemaVersion": 1, "sourceSha256": FROZEN_SOURCE_SHA256, "records": crosswalk}
        self.assertTrue(all(c.passed for c in validate_crosswalk(self.records, payload, FROZEN_SOURCE_SHA256)))
        changed = deepcopy(payload)
        changed["records"][0]["sourceRow"] = 999
        self.assertFalse(all(c.passed for c in validate_crosswalk(self.records, changed, FROZEN_SOURCE_SHA256)))
        self.assertFalse(all(c.passed for c in validate_crosswalk(self.records, payload, "changed")))

    def test_unresolved_siu_and_each_linked_conflict_block_core(self):
        resolved = self.resolved_review()
        self.assertTrue(all(c.passed for c in validate_siu_provenance(self.records, resolved, FROZEN_SOURCE_SHA256)))
        for key, value in (("status", "REQUIRES_AUTHOR_REVIEW"), ("mappingVerified", False),
                           ("linkedSourceMappingsVerified", False), ("potentialDuplicateResolved", False),
                           ("categoricalDisagreementsResolved", False), ("authorDecisionRequired", True),
                           ("questions", ["unresolved"]), ("membershipChangeProposed", True),
                           ("sourceSha256", "changed")):
            with self.subTest(key=key):
                changed = deepcopy(resolved)
                changed[key] = value
                self.assertFalse(all(c.passed for c in validate_siu_provenance(self.records, changed, FROZEN_SOURCE_SHA256)))

    def test_unknown_vital_status_cannot_be_reclassified(self):
        changed = deepcopy(self.records)
        unknown = next(r for r in changed if r["vital_status"] == "unknown")
        unknown["vital_status"] = "survival"
        self.assertFalse(self.check_named(validate_core_records(changed), "Documented vital status categories").passed)
        self.assertEqual(self.check_named(validate_core_records(changed), "Unknown vital status remains outside known denominator").observed, 38)

    def test_full_troponin_failures_are_preserved(self):
        full_records, ingestion, _, errors = load_source(DEFAULT_WORKBOOK)
        self.assertEqual(errors, [])
        self.assertTrue(all(c.passed for c in ingestion))
        full = validate_records(full_records)
        failures = [c for c in full if not c.passed]
        self.assertEqual(len(failures), 4)
        self.assertTrue(all("troponin" in c.name.lower() for c in failures))
        self.assertFalse(any("troponin" in c.name.lower() for c in self.checks))

    def test_executed_outputs_are_aggregate_only_and_preserve_private_inputs(self):
        inputs = [DEFAULT_WORKBOOK, DEFAULT_CROSSWALK, DEFAULT_SIU_REVIEW, DEFAULT_FULL_AUDIT]
        hashes = {path: sha256(path) for path in inputs}
        with tempfile.TemporaryDirectory() as directory:
            report, status = Path(directory) / "report.md", Path(directory) / "status.json"
            payload = run_core_validation(DEFAULT_WORKBOOK, report, status)
            self.assertEqual(payload["statuses"]["CORE_ATLAS_NUMERICAL_VALIDATION"], "PASS")
            self.assertEqual(payload["statuses"]["CORE_ATLAS_VALIDATION"], "FAIL")
            self.assertFalse(payload["patientExportGenerated"])
            def keys(value):
                if isinstance(value, dict):
                    return set(value) | set().union(*(keys(v) for v in value.values()))
                if isinstance(value, list):
                    return set().union(*(keys(v) for v in value))
                return set()
            serialized = json.loads(status.read_text())
            self.assertFalse(keys(serialized) & {"records", "sourceId", "sourceDoi", "sourceRow", "label", "evidence"})
            self.assertIn("CORE_ATLAS_VALIDATION", report.read_text())
        self.assertEqual({path: sha256(path) for path in inputs}, hashes)

    def test_changed_source_fingerprint_prevents_core_pass(self):
        with tempfile.TemporaryDirectory() as directory:
            report, status = Path(directory) / "report.md", Path(directory) / "status.json"
            with patch("core_validation.sha256", return_value="changed"):
                payload = run_core_validation(DEFAULT_WORKBOOK, report, status)
            self.assertFalse(next(c["passed"] for c in payload["checks"] if c["name"] == "Frozen authoritative source fingerprint"))
            self.assertEqual(payload["statuses"]["CORE_ATLAS_VALIDATION"], "FAIL")


if __name__ == "__main__":
    unittest.main()
