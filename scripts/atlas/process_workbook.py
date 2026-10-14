#!/usr/bin/env python3
"""Validate the unchanged authoritative workbook; gate a restricted local export.

Exit 0: every required frozen check passed and local review artifacts were made.
Exit 1: an invariant/target failed; no browser-facing patient artifact is retained.
This command never approves, commits, copies or publishes a patient dataset.
"""

import argparse
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import statistics
import subprocess
import tempfile
from typing import Any
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

import openpyxl
from scipy.stats import beta

from normalization import SourceError, normalize_row, review_records


REPOSITORY = Path(__file__).resolve().parents[2]
DEFAULT_WORKBOOK = REPOSITORY / "esearch_sources/Triple_M_Master_Patient_Dataset_Expanded_N38.xlsx"
OUTPUT_DIRECTORY = REPOSITORY / "esearch_sources/_processed"
CHECK_ONLY_AUDIT = OUTPUT_DIRECTORY / "core_atlas/full_study_validation.private.json"
MASTER_SHEET = "Master_Patient_Dataset_N38"
REQUIRED_HEADERS = (
    "Patient_ID", "DOI", "Age", "MG_Phenotype", "RNS_Performed", "RNS_Result",
    "SFEMG_Performed", "SFEMG_Result", "EMG_Conventional_Performed",
    "EMG_Interpretation", "Objective_NMJ_Category", "AChR_Performed",
    "AChR_Positive", "Troponin_Max", "LVEF_Reduced", "Outcome_Survival",
)


@dataclass(frozen=True)
class Check:
    name: str
    expected: Any
    observed: Any
    passed: bool
    dependency: str  # integrity, core categories, or non-exported descriptive
    note: str = ""


def exact_binomial_interval(successes: int, total: int) -> list[float] | None:
    """Clopper–Pearson, two-sided 95%; percentages rounded only for reporting."""
    if total <= 0 or not 0 <= successes <= total:
        return None
    lower = 0.0 if successes == 0 else float(beta.ppf(.025, successes, total - successes + 1))
    upper = 1.0 if successes == total else float(beta.ppf(.975, successes + 1, total - successes))
    return [round(lower * 100, 1), round(upper * 100, 1)]


def quantile(values: list[float], probability: float) -> float:
    """Linear interpolation (inclusive quartiles), matching the frozen age IQR."""
    ordered = sorted(values)
    position = (len(ordered) - 1) * probability
    low = int(position)
    fraction = position - low
    return ordered[low] + fraction * (ordered[min(low + 1, len(ordered) - 1)] - ordered[low])


def describe(values: list[float]) -> dict[str, Any]:
    if not values:
        return {"n": 0, "median": None, "iqr": None, "range": None}
    return {
        "n": len(values), "median": statistics.median(values),
        "iqr": [quantile(values, .25), quantile(values, .75)],
        "range": [min(values), max(values)],
    }


def load_source(workbook: Path) -> tuple[list[dict], list[Check], dict, list[str]]:
    """Read values and formulas separately; reject formulas in mapped input fields.

    data_only alone cannot prove the freshness of cached formula results. The
    inspected source has no formulas. A future mapped formula requires review.
    """
    checks: list[Check] = []
    details: dict = {}
    errors: list[str] = []
    if not workbook.is_file():
        return [], [Check("Authoritative workbook readable", True, False, False, "integrity", "Input file absent")], details, errors
    value_book = formula_book = None
    try:
        value_book = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
        formula_book = openpyxl.load_workbook(workbook, read_only=True, data_only=False)
        details["sheetNames"] = value_book.sheetnames
        has_sheet = MASTER_SHEET in value_book.sheetnames
        checks.append(Check("Authoritative master sheet", MASTER_SHEET, MASTER_SHEET if has_sheet else value_book.sheetnames, has_sheet, "integrity"))
        if not has_sheet:
            return [], checks, details, errors
        sheet = value_book[MASTER_SHEET]
        formula_sheet = formula_book[MASTER_SHEET]
        value_rows = list(sheet.iter_rows(values_only=True))
        formula_rows = list(formula_sheet.iter_rows())
        header = value_rows[0] if value_rows else ()
        details["headers"] = list(header)
        header_strings = [h for h in header if isinstance(h, str)]
        valid_headers = len(header_strings) == len(header) == len(set(header_strings))
        missing = sorted(set(REQUIRED_HEADERS) - set(header_strings))
        checks.append(Check("Unique complete source header mapping", [], missing or ([] if valid_headers else ["Invalid or duplicate header"]), valid_headers and not missing, "integrity"))
        if not valid_headers or missing:
            return [], checks, details, errors
        mapped_indexes = [header.index(name) for name in REQUIRED_HEADERS]
        formulas = [cell.coordinate for cells in formula_rows for index, cell in enumerate(cells)
                    if index in mapped_indexes and cell.data_type == "f"]
        errors_in_cells = [cell.coordinate for cells in formula_rows for index, cell in enumerate(cells)
                           if index in mapped_indexes and cell.data_type == "e"]
        checks.append(Check("Mapped formula cells (cache reliability)", 0, len(formulas), not formulas, "integrity", "Mapped formulas are rejected rather than trusting cached values"))
        checks.append(Check("Mapped spreadsheet error cells", 0, len(errors_in_cells), not errors_in_cells, "integrity"))
        details["formulaCoordinates"] = formulas
        details["errorCoordinates"] = errors_in_cells
        source_rows = [(number, dict(zip(header, cells))) for number, cells in enumerate(value_rows[1:], 2) if any(cell is not None for cell in cells)]
        details["sourceRowCount"] = len(source_rows)
        checks.append(Check("Source patient rows", 38, len(source_rows), len(source_rows) == 38, "integrity"))
        if formulas or errors_in_cells:
            return [], checks, details, errors
        normalized = []
        for number, row in source_rows:
            try:
                normalized.append(normalize_row(row, number))
            except (SourceError, TypeError, ValueError) as exc:
                # Record coordinates, never raw identifiers or narrative text.
                errors.append(f"Master sheet row {number}: {exc}")
        checks.append(Check("All source rows normalized without inferred categories", len(source_rows), len(normalized), not errors, "integrity"))
        checks.append(Check("No rows duplicated/dropped during processing", len(source_rows), len(normalized), len(source_rows) == len(normalized), "integrity"))
        return normalized, checks, details, errors
    except (OSError, ValueError, KeyError, BadZipFile, ParseError, openpyxl.utils.exceptions.InvalidFileException) as exc:
        checks.append(Check("Authoritative workbook readable", True, False, False, "integrity", f"{type(exc).__name__}: workbook could not be parsed reliably"))
        return [], checks, details, errors
    finally:
        if value_book is not None:
            value_book.close()
        if formula_book is not None:
            formula_book.close()


def validate_records(records: list[dict[str, Any]]) -> list[Check]:
    checks: list[Check] = []

    def check(name: str, expected: Any, observed: Any, dependency: str = "core categories", note: str = "") -> None:
        checks.append(Check(name, expected, observed, expected == observed, dependency, note))

    def counts(field: str, categories: tuple[str, ...]) -> dict:
        observed = Counter(row[field] for row in records)
        return {category: observed[category] for category in categories}

    check("Normalized patient records", 38, len(records), "integrity")
    check("Unique patient identifiers", 38, len({r["source_id"] for r in records}), "integrity")
    check("Unique source publications (DOI stripped/lowercased)", 27, len({r["source_doi"] for r in records}), "integrity")
    check("Source-to-record row mapping is one-to-one", True, len({r["source_row"] for r in records}) == len(records), "integrity")
    check("MG-like phenotype inclusion characteristic", 38, sum(r["mg_phenotype"] == "performed" for r in records), "integrity", "Cohort selection characteristic; no population prevalence denominator")
    achr_tested = [r for r in records if r["achr_status"] == "performed"]
    achr_positive = [r for r in records if r["achr_result"] == "positive"]
    achr_negative = [r for r in records if r["achr_result"] == "negative"]
    check("AChR tested / positive / negative / not documented", [30, 11, 19, 8], [len(achr_tested), len(achr_positive), len(achr_negative), len(records) - len(achr_tested)])
    check("AChR performed/result consistency", 0, sum((r["achr_status"] == "performed") != (r["achr_result"] in ("positive", "negative")) for r in records))
    check("AChR positivity (%)", 36.7, round(len(achr_positive) / len(achr_tested) * 100, 1) if achr_tested else None)
    check("AChR exact 95% CI (%)", [19.9, 56.1], exact_binomial_interval(len(achr_positive), len(achr_tested)))
    formal = [r for r in records if r["formal_nmj"]]
    rns = [r for r in records if r["rns_status"] == "performed"]
    sfemg = [r for r in records if r["sfemg_status"] == "performed"]
    check("Formal RNS union SFEMG patients", 4, len(formal))
    check("RNS patients", 4, len(rns))
    check("SFEMG patients", 1, len(sfemg))
    check("SFEMG contained in RNS subset (no fifth patient)", True, {r["source_id"] for r in sfemg} <= {r["source_id"] for r in rns})
    check("Formal test availability equals source union", True, all(r["formal_nmj"] == (r["rns_status"] == "performed" or r["sfemg_status"] == "performed") for r in records))
    check("Performed RNS/SFEMG results explicitly negative", True, all(r["rns_result"] == "negative" for r in rns) and all(r["sfemg_result"] == "negative" for r in sfemg))
    abnormal = sum(any(r[f"{test}_status"] == "performed" and r[f"{test}_result"] == "abnormal" for test in ("rns", "sfemg")) for r in formal)
    check("Objective abnormality among formally evaluated", [0, 4], [abnormal, len(formal)])
    check("Objective NMJ exact 95% CI (%)", [0.0, 60.2], exact_binomial_interval(abnormal, len(formal)))
    check("Objective NMJ evidence categories", {"formal_negative": 4, "indeterminate_qualitative": 2, "not_documented": 32}, counts("nmj_evidence", ("formal_negative", "indeterminate_qualitative", "not_documented")))
    check("Formal-negative category requires a documented formal negative test", True, all((r["nmj_evidence"] == "formal_negative") == r["formal_nmj"] for r in records))
    check("Qualitative indeterminate evidence lacks formal confirmation", 0, sum(r["nmj_evidence"] == "indeterminate_qualitative" and r["formal_nmj"] for r in records))
    check("AChR-positive / formal NMJ overlap", 0, sum(r["formal_nmj"] for r in achr_positive), note="Critical cross-variable condition: absence of paired evaluation is not a negative test")
    needle = [r for r in records if r["needle_status"] == "performed"]
    check("Conventional needle EMG evaluated / not documented", [11, 27], [len(needle), len(records) - len(needle)])
    check("Needle EMG categories among evaluated", {"myopathic": 6, "neuropathic": 1, "normal": 1, "indeterminate": 3}, {category: sum(r["needle_emg"] == category for r in needle) for category in ("myopathic", "neuropathic", "normal", "indeterminate")})
    check("Needle EMG performed/interpretation consistency", 0, sum((r["needle_status"] == "performed") != (r["needle_emg"] in ("myopathic", "neuropathic", "normal", "indeterminate")) for r in records))
    myopathic = sum(r["needle_emg"] == "myopathic" for r in needle)
    check("Myopathic needle EMG (%)", 54.5, round(myopathic / len(needle) * 100, 1) if needle else None)
    check("Needle EMG exact 95% CI (%) (audit)", [23.4, 83.3], exact_binomial_interval(myopathic, len(needle)))
    check("Vital status categories", {"death": 16, "survival": 21, "unknown": 1}, counts("vital_status", ("death", "survival", "unknown")))
    deaths = sum(r["vital_status"] == "death" for r in records)
    known = sum(r["vital_status"] != "unknown" for r in records)
    check("Mortality all patients (%)", 42.1, round(deaths / len(records) * 100, 1) if records else None)
    check("Mortality known vital status denominator", 37, known)
    check("Mortality known vital status (%)", 43.2, round(deaths / known * 100, 1) if known else None)
    check("Mortality all exact 95% CI (%) (audit)", [26.3, 59.2], exact_binomial_interval(deaths, len(records)))
    ages = describe([r["age"] for r in records if r["age"] is not None])
    check("Strictly numeric ages n", 28, ages["n"], "non-exported descriptive")
    check("Age median (years)", 71.0, ages["median"], "non-exported descriptive")
    check("Age IQR (years)", [66.5, 78.0], ages["iqr"], "non-exported descriptive")
    check("Age range (years, audit)", [42.0, 86.0], ages["range"], "non-exported descriptive")
    troponins = describe([r["troponin"].value_ng_l for r in records if r["troponin"].status == "quantitative"])
    check("Supported quantitative peak troponin n", 33, troponins["n"], "non-exported descriptive", "Exclude censored, interval, unitless-peak and competing numeric-assay entries; preserve literal unit conversions")
    check("Troponin median (ng/L)", 1506.0, troponins["median"], "non-exported descriptive")
    check("Troponin IQR (ng/L)", [1028.0, 3891.0], troponins["iqr"], "non-exported descriptive")
    check("Troponin range (ng/L, audit)", [498.0, 162000.0], troponins["range"], "non-exported descriptive", "Source ng/mL cannot silently become ng/L")
    check("Categorical LVEF", {"preserved": 21, "reduced": 5, "not_documented": 12}, counts("lvef", ("preserved", "reduced", "not_documented")), "non-exported descriptive", "Use source LVEF_Reduced classification, including pre-existing dysfunction; do not infer new ICI injury")
    return checks


def private_directory(directory: Path) -> None:
    """The sole patient artifact destination must be ignored and outside deployment."""
    if directory.resolve() != OUTPUT_DIRECTORY.resolve():
        raise ValueError("Patient artifacts must use the repository's ignored esearch_sources/_processed directory")
    result = subprocess.run(["git", "check-ignore", "-q", str(directory / "atlas.review.json")], cwd=REPOSITORY, check=False)
    if result.returncode != 0:
        raise ValueError("Restricted output directory is not ignored by Git")
    tracked = subprocess.run(["git", "ls-files", "--", str(directory)], cwd=REPOSITORY, text=True, capture_output=True, check=True)
    if tracked.stdout.strip():
        raise ValueError("Restricted output directory contains tracked files")
    directory.mkdir(parents=True, exist_ok=True)
    directory.chmod(0o700)


def atomic_json(path: Path, payload: Any) -> None:
    with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=path.parent, prefix=f".{path.name}.", delete=False) as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2, allow_nan=False)
        handle.write("\n")
        temporary = Path(handle.name)
    temporary.chmod(0o600)
    temporary.replace(path)


def render_report(checks: list[Check], errors: list[str], records: list[dict], export_written: bool, timestamp: str, check_only: bool = False) -> str:
    passed = bool(checks) and all(check.passed for check in checks)
    failed = [check for check in checks if not check.passed]
    core = [check for check in checks if check.dependency != "non-exported descriptive"]
    core_passed = bool(core) and all(check.passed for check in core)

    def table_value(value: Any) -> str:
        return json.dumps(value, ensure_ascii=False).replace("|", "\\|")

    lines = [
        "# Phase 3A data validation report", "", f"Executed at: {timestamp}.", "",
        f"**Overall frozen validation: {'PASS' if passed else 'FAIL — publication gate closed'}.**",
        f"{'Legacy frozen categorical/workbook checks' if check_only else 'Core categorical and integrity checks'}: {'PASS' if core_passed else 'FAIL'}. "
        f"Required checks: {sum(c.passed for c in checks)}/{len(checks)} passed.", "",
        "The authoritative workbook is now present in the documented `esearch_sources/` equivalent. "
        "This replaces the former workbook-availability blocker with executed ingestion and validation. "
        "The original workbook was opened read-only and was not corrected, imputed or rewritten.", "",
        "## Executed checks", "",
        "Expected values are independent frozen targets, never replacement patient records. "
        "Confidence intervals use two-sided exact Clopper–Pearson calculations (`scipy.stats.beta`). "
        "Quartiles use inclusive linear interpolation. Percentages and interval bounds are rounded to one decimal for comparison.", "",
        "| Check | Export dependency | Expected | Observed | Status |",
        "| --- | --- | --- | --- | --- |",
    ]
    lines.extend(f"| {check.name} | {check.dependency} | {table_value(check.expected)} | {table_value(check.observed)} | {'PASS' if check.passed else 'FAIL'} |" for check in checks)
    lines += ["", "Integrity and core-category checks support the six minimal browser record fields. "
              "Age, troponin and LVEF are explicitly non-exported descriptive validations. "
              "The current pipeline requires every frozen check to pass before producing its local review export; "
              "a passing core subset alone does not declare the full study validated.", "", "## Discrepancies and source ambiguities", ""]
    if failed:
        lines += [f"- {c.name}: expected {table_value(c.expected)}; observed {table_value(c.observed)}. {c.note}" for c in failed]
    else:
        lines.append("No required frozen target differed under the documented normalization rules.")
    if errors:
        lines += ["", "Source mapping errors (row coordinates only; no raw identifiers):", ""] + [f"- {error}" for error in errors]
    if records:
        lines += ["", "Observed AChR/formal-NMJ cross-tabulation (source records, not aggregate reconstruction):", "",
                  "| AChR category | Formal testing documented | Formal testing not documented | Total |",
                  "| --- | ---: | ---: | ---: |"]
        for category, display in (("positive", "Positive"), ("negative", "Negative"), ("not_documented", "Not documented")):
            members = [r for r in records if (r["achr_result"] if r["achr_result"] in ("positive", "negative") else "not_documented") == category]
            documented = sum(r["formal_nmj"] for r in members)
            lines.append(f"| {display} | {documented} | {len(members) - documented} | {len(members)} |")
        troponin_status = Counter(r["troponin"].status for r in records)
        reasons = Counter(r["troponin"].reason for r in records if r["troponin"].status == "ambiguous")
        lines += ["", f"Troponin parsing status: {table_value(dict(sorted(troponin_status.items())))}.", ""]
        lines += [f"- {reason}: {count} source entry/entries." for reason, count in sorted(reasons.items())]
        literal_ng_ml = [r["troponin"].value_ng_l for r in records if "ng/mL" in r["troponin_raw_value"] and r["troponin"].value_ng_l is not None]
        if literal_ng_ml:
            lines += ["", f"Source concentrations explicitly labeled ng/mL convert literally to a maximum of {max(literal_ng_ml):,.0f} ng/L. "
                      "The audit's frozen upper range is 162,000 ng/L, checked separately above. "
                      "The pipeline retains the source unit and never silently changes ng/mL to ng/L."]
        if reasons:
            lines += ["", "Source intervals, lower bounds, unitless peak text and rows giving quantitative T and I assay "
                      "results lack a sufficient exact peak/assay selection rule. The audit does not document how it resolved "
                      "these entries. Clinical/source-author adjudication is required."]
        lines += ["",
                  "The audit also records a secondary troponin mean/SD difference from its manuscript "
                  "(8,305.78±28,037.83 versus audited 8,349.51±28,028.10). This pipeline does not reconstruct "
                  "or validate a mixed-assay mean by adopting those undocumented extraction choices.", "",
                  "Other preserved ambiguities: explicit non-performance differs from non-reporting internally; "
                  "the browser safely coarsens both to not documented. The mixed myopathic/polyneuropathic "
                  "needle case follows the frozen source's primary myopathic interpretation. LVEF categories "
                  "include pre-existing dysfunction and do not establish ICI causality. Vital status follows the "
                  "documented outcome prefix, never a narrative inference (including discharge/hospice wording)."]
    lines += ["", "## Publication boundary", "",
              "A private audit artifact and, where source IDs are unique, a provenance crosswalk are retained only "
              "in ignored `esearch_sources/_processed/`. They contain no additional published approval. "
              "Raw identifiers, source DOI, workbook metadata and record provenance are never bundled into the UI.", "",
              ("Read-only revalidation wrote aggregate review outputs only. Existing private audit, crosswalk and source-adjudication records were preserved. No patient artifact was generated or revoked by this invocation."
               if check_only else "The minimal `atlas.review.json` was written for local review only. It is not approved for publication."
               if export_written else "No browser-facing patient dataset was generated. Any stale `atlas.review.json` is removed on failure. "
               "The patient matrix must remain behind the data-validation/publication gate; aggregate context may still be displayed."), "",
              "No source correction, publication approval, Git commit, push or deployment was performed. "
              "See `docs/atlas-data-review.md` for the explicit review process and rerun instructions.", "",
              "## Scope of this report", "",
              "This report records only the executed Python ingestion and numerical checks above. "
              "Python regression tests, frontend build/lint/browser checks and production exclusion checks "
              "are reported separately by the delivery workflow; their success is not implied here."]
    return "\n".join(line.rstrip() for line in lines) + "\n"


def check_only(workbook: Path, report: Path) -> tuple[bool, list[Check]]:
    """Execute the unchanged full-study rules without replacing prior artifacts.

    The only private output is a NEW aggregate audit in the Phase 3A.2 folder.
    This mode has no export, crosswalk or approval operation.
    """
    if workbook.resolve() in (report.resolve(), CHECK_ONLY_AUDIT.resolve()):
        raise ValueError("Read-only validation outputs cannot overwrite their workbook input")
    if report.resolve().is_relative_to((REPOSITORY / "esearch_sources").resolve()) or report.resolve().is_relative_to((REPOSITORY / "research_sources").resolve()):
        raise ValueError("A read-only aggregate report cannot overwrite research artifacts")
    if report.resolve() == (REPOSITORY / "DATA_VALIDATION_REPORT.md").resolve():
        raise ValueError("Read-only revalidation preserves the historical full-study report")
    if CHECK_ONLY_AUDIT.parent.resolve() != (OUTPUT_DIRECTORY / "core_atlas").resolve():
        raise ValueError("Read-only audit must remain in the restricted Phase 3A.2 directory")
    ignored = subprocess.run(["git", "check-ignore", "-q", str(CHECK_ONLY_AUDIT)], cwd=REPOSITORY, check=False)
    tracked = subprocess.run(["git", "ls-files", "--", str(CHECK_ONLY_AUDIT)], cwd=REPOSITORY, text=True, capture_output=True, check=True)
    if ignored.returncode != 0 or tracked.stdout.strip():
        raise ValueError("Read-only audit destination must be ignored and untracked")
    timestamp = datetime.now(timezone.utc).isoformat(timespec="seconds")
    try:
        source_hash = hashlib.sha256(workbook.read_bytes()).hexdigest()
    except OSError:
        source_hash = None
    records, checks, _, errors = load_source(workbook)
    if records:
        checks += validate_records(records)
    try:
        unchanged = source_hash is not None and hashlib.sha256(workbook.read_bytes()).hexdigest() == source_hash
    except OSError:
        unchanged = False
    checks.append(Check("Original workbook unchanged during ingestion", True, unchanged, unchanged, "integrity"))
    passed = bool(checks) and all(check.passed for check in checks)
    CHECK_ONLY_AUDIT.parent.mkdir(parents=True, exist_ok=True)
    CHECK_ONLY_AUDIT.parent.chmod(0o700)
    atomic_json(CHECK_ONLY_AUDIT, {
        "schemaVersion": 1, "mode": "CHECK_ONLY", "sourceSha256": source_hash,
        "validatedAt": timestamp, "validationPassed": passed,
        "checks": [asdict(check) for check in checks], "normalizationErrorCount": len(errors),
        "patientExportGenerated": False, "originalAdjudicationRecordsRewritten": False,
    })
    report.parent.mkdir(parents=True, exist_ok=True)
    text = render_report(checks, errors, records, False, timestamp, check_only=True)
    text = text.replace("# Phase 3A data validation report", "# Full-study read-only revalidation — Phase 3A.2", 1)
    text += "\nLegacy numerical agreement is separate from the independent core-atlas source-integrity decision. " \
        "This full-study report does not establish CORE_ATLAS_VALIDATION or resolve linked provenance questions; " \
        "see [the separately executed scoped report](CORE_ATLAS_VALIDATION_REPORT.md).\n"
    report.write_text(text, encoding="utf-8")
    return passed, checks


def run(workbook: Path, report: Path) -> tuple[bool, list[Check]]:
    private_directory(OUTPUT_DIRECTORY)
    export_path = OUTPUT_DIRECTORY / "atlas.review.json"
    # Invalidate review data before every new attempt, including hashing/read
    # failures. A previous successful source must not survive a failed rerun.
    export_path.unlink(missing_ok=True)
    timestamp = datetime.now(timezone.utc).isoformat(timespec="seconds")
    try:
        source_hash = hashlib.sha256(workbook.read_bytes()).hexdigest() if workbook.is_file() else None
    except OSError:
        source_hash = None
    records, checks, details, errors = load_source(workbook)
    if records:
        checks += validate_records(records)
    try:
        unchanged = source_hash is not None and workbook.is_file() and hashlib.sha256(workbook.read_bytes()).hexdigest() == source_hash
    except OSError:
        unchanged = False
    checks.append(Check("Original workbook unchanged during ingestion", True, unchanged, unchanged, "integrity"))
    passed = bool(checks) and all(check.passed for check in checks)
    public, provenance = [], []
    if records and len({r["source_id"] for r in records}) == len(records):
        public, provenance = review_records(records)
        atomic_json(OUTPUT_DIRECTORY / "atlas.provenance.private.json", {
            "schemaVersion": 1, "sourceSha256": source_hash, "validatedAt": timestamp,
            "validationPassed": passed, "records": provenance,
        })
    else:
        (OUTPUT_DIRECTORY / "atlas.provenance.private.json").unlink(missing_ok=True)
    atomic_json(OUTPUT_DIRECTORY / "validation.audit.private.json", {
        "validationPassed": passed, "sourceSha256": source_hash, "validatedAt": timestamp,
        "workbookInspection": details, "normalizationErrors": errors,
        "checks": [asdict(check) for check in checks],
        "availabilityCounts": {field: dict(Counter(r[field] for r in records)) for field in ("rns_status", "sfemg_status", "achr_status", "needle_status")},
        "troponinAdjudications": [{
            "sourceRow": r["source_row"], "sourceId": r["source_id"],
            "rawValue": r["troponin_raw_value"], "status": r["troponin"].status,
            "reason": r["troponin"].reason or "Literal ng/mL conversion requires reconciliation with the frozen audit",
            "normalizedValueNgL": r["troponin"].value_ng_l,
        } for r in records if r["troponin"].status == "ambiguous" or "ng/mL" in r["troponin_raw_value"]],
    })
    if passed:
        atomic_json(export_path, {"schemaVersion": 1, "validation": {"passed": True, "sourceSha256": source_hash, "validatedAt": timestamp}, "records": public})
    else:
        # Revoke a stale local review artifact; never leave previously passed
        # patient data available after a changed source fails a required check.
        export_path.unlink(missing_ok=True)
    report.parent.mkdir(parents=True, exist_ok=True)
    report.write_text(render_report(checks, errors, records, passed, timestamp), encoding="utf-8")
    return passed, checks


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--report", type=Path)
    parser.add_argument("--check-only", action="store_true", help="Preserve prior private artifacts; write new aggregate review outputs only")
    arguments = parser.parse_args()
    report = arguments.report or REPOSITORY / ("docs/FULL_STUDY_REVALIDATION_REPORT.md" if arguments.check_only else "DATA_VALIDATION_REPORT.md")
    try:
        validator = check_only if arguments.check_only else run
        passed, checks = validator(arguments.workbook.resolve(), report.resolve())
    except (OSError, ValueError, subprocess.SubprocessError) as exc:
        # A report/output error after validation must also revoke review data.
        if not arguments.check_only:
            try:
                (OUTPUT_DIRECTORY / "atlas.review.json").unlink(missing_ok=True)
            except OSError:
                pass
        parser.exit(1, f"Validation blocked: {exc}\n")
    print(f"{'PASS' if passed else 'FAIL'}: {sum(c.passed for c in checks)}/{len(checks)} required checks; report: {report}")
    for check in checks:
        if not check.passed:
            print(f"  {check.name}: expected {check.expected!r}; observed {check.observed!r}")
    print("Read-only full-study checks; previous artifacts unchanged and no patient export." if arguments.check_only else "Local review export written; publication remains unapproved." if passed else "No browser-facing patient export. Source review required.")
    return 0 if passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
